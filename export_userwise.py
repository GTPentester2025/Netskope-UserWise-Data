#!/usr/bin/env python3
"""
User-wise Netskope events for your own list of apps and of websites.

Two list files next to this script, one entry per line, # for comments
(both git-ignored, so the lists never leave this machine except as the
API filter). A missing or empty file just skips that report.
    apps.txt       app names exactly as Netskope shows them (e.g. ChatGPT),
                   matched against `app`
    websites.txt   site names / domains / hostnames as in the Netskope
                   "Website" column (e.g. sharecad), matched against
                   `site`, `domain` and `hostname`

Output, one Excel file per list, one sheet, one row per event, sorted by
user then time, rows identical in every column dropped. Traffic and activity events share the sheet; the Event
Type column tells them apart:
    output/Apps_<from>_to_<to>.xlsx
    output/Websites_<from>_to_<to>.xlsx
        Traffic    /api/v2/events/datasearch/page
                   browser, source IP, OS, device, bytes
        Activity   /api/v2/events/datasearch/application
                   activity, object, file path / size
A sheet past Excel's row limit continues on "Events 2", "Events 3", ...

The HTTP, retry, paging and time-window logic is imported from
../Netskope-API/export_genai_page_bytes.py, same as bulk_upload_uba.py.

Usage:
    python export_userwise.py 2026-08-01 2026-09-30   -> both days inclusive
    python export_userwise.py --selftest              -> offline check
"""

import csv
import os
import sys
import time
from datetime import datetime, timedelta
from typing import Any, Dict, Iterator, List, Tuple

import pandas as pd
from openpyxl import Workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(SCRIPT_DIR, "..", "Netskope-API"))
import export_genai_page_bytes as ns  # noqa: E402

# ---------------------------------------------------------------------------
# CONFIGURATION
# ---------------------------------------------------------------------------
# SECURITY: the token lives in this file -- keep it out of git / shares.
API_TOKEN: str = "PUT_YOUR_API_TOKEN_HERE"
BASE_URL: str = "https://YOUR_TENANT.goskope.com"
AUTH_MODE: str = "bearer"                 # flip to "netskope" on 401s

# Event Type label -> endpoint. Drop one to skip that kind of event.
ENDPOINTS: Dict[str, str] = {
    "Traffic": "/api/v2/events/datasearch/page",
    "Activity": "/api/v2/events/datasearch/application",
}

# API field -> Excel header, in output order. The first twelve match the
# Netskope Skope IT traffic export. Fields an event lacks export blank.
COLUMNS: Dict[str, str] = {
    "timestamp": "Timestamp (UTC) Seconds",   # raw epoch seconds
    "browser": "Browser",
    "user": "User",
    "srcip": "Source IP",
    "os": "OS",
    "device": "Device",
    "usergroup": "User Group",
    "organization_unit": "Organization Unit",
    "website": "Website",          # first non-blank of site/hostname/domain
    "numbytes": "Total Bytes",
    "client_bytes": "Bytes Uploaded",
    "server_bytes": "Bytes Downloaded",
    "event_type": "Event Type",
    "app": "Application",
    "url": "URL",
    "activity": "Activity",
    "object_id": "Object ID",
    "object": "Object Name",
    "object_type": "Object Type",
    "file_path": "File Path",
    "file_size": "File Size (Bytes)",
}
WEBSITE_FIELDS: List[str] = ["site", "hostname", "domain"]

# Headers written as numbers in Excel.
NUMERIC_HEADERS = {"Timestamp (UTC) Seconds", "Total Bytes", "Bytes Uploaded",
                   "Bytes Downloaded", "File Size (Bytes)"}

# (report, list file, fields each entry is matched against)
REPORTS = [
    ("Apps", os.path.join(SCRIPT_DIR, "apps.txt"), ["app"]),
    ("Websites", os.path.join(SCRIPT_DIR, "websites.txt"), WEBSITE_FIELDS),
]

# Entries per query, keeps the URL a sane length.
TARGET_BATCH: int = 50

SHEET_TITLE: str = "Events"
# Excel's 1,048,576 rows minus the header.
SHEET_ROWS: int = 1_048_575

OUTPUT_DIR: str = os.path.join(SCRIPT_DIR, "output")
LOG_DIR: str = os.path.join(SCRIPT_DIR, "logs")

# ---------------------------------------------------------------------------


def read_targets(path: str) -> List[str]:
    """Non-blank, non-# lines, de-duplicated, order kept. Missing file -> []."""
    if not os.path.exists(path):
        return []
    with open(path, encoding="utf-8-sig") as fh:
        lines = (line.strip() for line in fh)
        return list(dict.fromkeys(l for l in lines if l and not l.startswith("#")))


def build_query(fields: List[str], targets: List[str]) -> str:
    """app in [...] or site in [...] -- quoted here, not via ns.build_term,
    so names containing commas survive."""
    values = "[" + ", ".join(ns.quote_value(t) for t in targets) + "]"
    return " or ".join(f"{field} in {values}" for field in fields)


def day_windows(first: str, last: str) -> List[Tuple[int, int]]:
    """One (local midnight, 23:59:59) pair per day, `first`..`last` inclusive.
    Fetching and writing a day at a time keeps memory to one day's events."""
    day = datetime.strptime(first, "%Y-%m-%d")
    stop = datetime.strptime(last, "%Y-%m-%d")
    if stop < day:
        raise ValueError("end date is before start date")
    windows = []
    while day <= stop:
        nxt = day + timedelta(days=1)
        windows.append((int(day.timestamp()), int(nxt.timestamp()) - 1))
        day = nxt
    return windows


def to_row(record: Dict[str, Any], event_type: str) -> Dict[str, Any]:
    """Flatten, copy byte-field aliases onto numbytes/client_bytes/
    server_bytes, fill `website` and tag the event type."""
    row = ns.flatten_record(record)
    for canonical, aliases in ns.FIELD_ALIASES.items():
        if row.get(canonical) in (None, ""):
            row[canonical] = next((row[a] for a in aliases
                                   if row.get(a) not in (None, "")), "")
    row["website"] = next((row[f] for f in WEBSITE_FIELDS if row.get(f)), "")
    row["event_type"] = event_type
    return row


def to_frame(rows: List[Dict[str, Any]]) -> pd.DataFrame:
    frame = pd.DataFrame(rows).reindex(columns=list(COLUMNS))
    return frame.rename(columns=COLUMNS).fillna("")


def to_number(value: str) -> Any:
    if value == "":
        return None                # empty cell, not a text ""
    for kind in (int, float):
        try:
            return kind(value)
        except ValueError:
            pass
    return value


def csv_rows(path: str) -> Iterator[List[str]]:
    with open(path, newline="", encoding="utf-8") as fh:
        yield from csv.reader(fh)


def add_sheets(book: Workbook, title: str, csv_path: str) -> int:
    """Copy one CSV into the workbook without duplicate rows, sorted by user
    then time, rolling over to "<title> 2", ... at the Excel row limit.
    Returns the row count."""
    try:
        frame = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
        # Identical in every column = duplicate (e.g. one event matched by
        # both its site and its hostname batch, or Netskope repeating a row).
        before = len(frame)
        frame = frame.drop_duplicates()
        ns.log.info("%s %d duplicate rows dropped", title, before - len(frame))
        frame["_ts"] = pd.to_numeric(frame[COLUMNS["timestamp"]], errors="coerce")
        frame = frame.sort_values([COLUMNS["user"], "_ts"], kind="stable")
        frame = frame.drop(columns="_ts")
        header, rows = list(frame.columns), frame.itertuples(index=False, name=None)
    except MemoryError:
        ns.log.warning("%s too big to sort/dedupe in memory, written in day "
                       "order WITH duplicates", title)
        rows = csv_rows(csv_path)
        header = next(rows)
    numeric = [h in NUMERIC_HEADERS for h in header]

    sheet, count = None, 0
    for row in rows:
        if count % SHEET_ROWS == 0:
            part = count // SHEET_ROWS + 1
            sheet = book.create_sheet(title if part == 1 else f"{title} {part}")
            sheet.append(header)
        sheet.append([to_number(v) if n else v for v, n in zip(row, numeric)])
        count += 1
    if sheet is None:
        book.create_sheet(title).append(header)
    return count


def fetch_to_csv(session, stats: Dict[str, int], query: str, event_type: str,
                 windows: List[Tuple[int, int]], path: str, fresh: bool) -> int:
    """Pull one endpoint's events a day at a time, appending each day to
    `path`. `fresh` truncates the file and writes the header on day one."""
    ns.ENDPOINT_PATH = ENDPOINTS[event_type]
    total = 0
    for start, end in windows:
        rows = [to_row(r, event_type)
                for r in ns.fetch_all(session, query, start, end, stats)]
        to_frame(rows).to_csv(path, index=False, encoding="utf-8",
                              mode="w" if fresh else "a", header=fresh)
        fresh = False
        total += len(rows)
        ns.log.info("%s %-8s %d events (total %d)",
                    time.strftime("%Y-%m-%d", time.localtime(start)),
                    event_type, len(rows), total)
    return total


def main() -> int:
    for name in ("API_TOKEN", "BASE_URL", "AUTH_MODE"):
        setattr(ns, name, globals()[name])
    ns.LOG_DIR, ns.LOG_BASENAME = LOG_DIR, "export_userwise.log"
    # Ask only for the exported fields, not the ~100 Netskope returns by
    # default. Unknown names are ignored by the API.
    ns.SEND_FIELDS_PARAM = True
    ns.PLACEHOLDER_COLUMNS = set()
    ns.DESIRED_COLUMNS = ([c for c in COLUMNS if c not in ("website", "event_type")]
                          + WEBSITE_FIELDS)
    ns.setup_logging(time.strftime("%H%M%S"))

    if "PUT_YOUR_API_TOKEN" in API_TOKEN or "YOUR_TENANT" in BASE_URL:
        ns.log.error("CONFIG ERROR: set API_TOKEN / BASE_URL at the top of %s",
                     os.path.basename(__file__))
        return 1
    try:
        windows = day_windows(sys.argv[1], sys.argv[2])
    except (IndexError, ValueError) as exc:
        ns.log.error("BAD DATES: %s (use: export_userwise.py 2026-08-01 2026-09-30)", exc)
        return 1
    label = f"{sys.argv[1]}_to_{sys.argv[2]}"
    lists = {report: read_targets(path) for report, path, _ in REPORTS}
    if not any(lists.values()):
        ns.log.error("apps.txt and websites.txt are both missing or empty")
        return 1

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    session = ns.build_session()
    stats: Dict[str, int] = {"calls": 0}
    grand_total = 0
    ns.log.info("range      %s (%d days)", label, len(windows))

    for report, _, fields in REPORTS:
        targets = lists[report]
        if not targets:
            ns.log.info("%-10s skipped (no list)", report)
            continue
        ns.log.info("%-10s %d entries", report, len(targets))
        # Working file; removed once the Excel file is saved.
        csv_path = os.path.join(OUTPUT_DIR, f"{report}_{label}.csv")
        fresh = True
        try:
            for event_type in ENDPOINTS:
                for i in range(0, len(targets), TARGET_BATCH):
                    query = build_query(fields, targets[i:i + TARGET_BATCH])
                    fetch_to_csv(session, stats, query, event_type, windows,
                                 csv_path, fresh)
                    fresh = False
        except ns.ApiError as exc:
            ns.log.error("NETSKOPE FAILED: %s (days so far are in %s)", exc, csv_path)
            return 2

        book = Workbook(write_only=True)
        count = add_sheets(book, SHEET_TITLE, csv_path)
        xlsx = os.path.join(OUTPUT_DIR, f"{report}_{label}.xlsx")
        book.save(xlsx)
        os.remove(csv_path)
        grand_total += count
        ns.log.info("wrote      %s (%d rows)", xlsx, count)

    ns.log.info("API calls  %d", stats["calls"])
    return 0 if grand_total else 3


def selftest() -> None:
    import tempfile

    assert build_query(["app"], ["ChatGPT", "a, b"]) == "app in ['ChatGPT', 'a, b']"
    assert build_query(WEBSITE_FIELDS, ["x.com"]) == (
        "site in ['x.com'] or hostname in ['x.com'] or domain in ['x.com']")

    windows = day_windows("2026-08-01", "2026-09-30")
    assert len(windows) == 61
    assert time.localtime(windows[0][0])[:6] == (2026, 8, 1, 0, 0, 0)
    assert time.localtime(windows[-1][1])[:6] == (2026, 9, 30, 23, 59, 59)
    assert all(b[0] == a[1] + 1 for a, b in zip(windows, windows[1:]))  # no gaps
    try:
        day_windows("2026-09-30", "2026-08-01")
        raise AssertionError("reversed range accepted")
    except ValueError:
        pass

    row = to_row({"user": "a@x.com", "hostname": "chat.x.com", "domain": "x.com",
                  "client_bytes_total": 5, "numbytes": 9}, "Traffic")
    assert row["website"] == "chat.x.com" and row["event_type"] == "Traffic"
    assert (row["client_bytes"], row["numbytes"], row["server_bytes"]) == (5, 9, "")
    frame = to_frame([row])
    assert list(frame.columns) == list(COLUMNS.values())
    assert list(frame.columns)[:12] == [
        "Timestamp (UTC) Seconds", "Browser", "User", "Source IP", "OS", "Device",
        "User Group", "Organization Unit", "Website", "Total Bytes",
        "Bytes Uploaded", "Bytes Downloaded"]          # same as Skope IT export

    # Sort by user then time (numeric, not text), rollover, numbers.
    global SHEET_ROWS
    saved, SHEET_ROWS = SHEET_ROWS, 2
    try:
        path = os.path.join(tempfile.mkdtemp(), "t.csv")
        pd.DataFrame({"User": ["b", "a", "a", "b"],
                      "Timestamp (UTC) Seconds": ["5", "10", "9", "5"],
                      "Total Bytes": ["10", "", "3", "10"]}).to_csv(path, index=False)
        book = Workbook()
        assert add_sheets(book, "Events", path) == 3
        assert book.sheetnames == ["Sheet", "Events", "Events 2"]
        assert list(book["Events"].values)[1:] == [("a", 9, 3), ("a", 10, None)]
        assert list(book["Events 2"].values)[1:] == [("b", 5, 10)]
    finally:
        SHEET_ROWS = saved
    print("selftest OK")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        sys.exit(0)
    sys.exit(main())
